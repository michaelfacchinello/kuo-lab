#!/usr/bin/env python3
"""
Convert the filled-in Kuo-Lab-Content-Workbook.xlsx into website content files.

Usage:
    python3 content-forms/import-workbook.py content-forms/Kuo-Lab-Content-Workbook.xlsx

Reads the Publications, Alumni, News, and Projects sheets and writes one
Markdown file per row into the matching src/content/ folder. The example row
(the yellow one) and blank rows are skipped automatically. Existing files are
left alone unless --overwrite is passed.

This does NOT touch People — member profiles come in on the Word form and are
short enough to add by hand (or ask Claude to do it).
"""
import sys
import re
import argparse
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "src" / "content"

# The example ("yellow") row is identified by the value in ONE specific column
# per sheet — never by scanning every cell, or a real row that happens to reuse
# an example value (e.g. a publication in the "Chromatin & Epigenetics" project)
# would be dropped by mistake.
EXAMPLE_BY_SHEET = {
    "Publications": ("Title", "Regulation of chromatin state during cell division"),
    "Alumni": ("Name", "Mary Major"),
    "News": ("Headline", "Kuo Lab paper accepted at JCB"),
    "Projects": ("Title", "Chromatin & Epigenetics"),
}


def slugify(text):
    text = (text or "").strip().lower()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_]+", "-", text)
    return re.sub(r"-+", "-", text).strip("-") or "untitled"


def q(value):
    """YAML single-quoted scalar; doubles internal single-quotes."""
    return "'" + str(value).strip().replace("'", "''") + "'"


def yaml_list(raw):
    items = [x.strip() for x in str(raw or "").split(",") if x.strip()]
    return "[" + ", ".join(q(i) for i in items) + "]"


def rows(ws, header_row=5, first_data=8):
    headers = [ws.cell(header_row, i).value for i in range(1, ws.max_column + 1)]
    for r in range(first_data, ws.max_row + 1):
        values = [ws.cell(r, i).value for i in range(1, ws.max_column + 1)]
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        record = {h: (v if v is not None else "") for h, v in zip(headers, values)}
        yield record


def write(path, body, overwrite):
    if path.exists() and not overwrite:
        print(f"  skip (exists): {path.relative_to(ROOT)}")
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding="utf-8")
    print(f"  wrote: {path.relative_to(ROOT)}")
    return 1


def is_example(record, sheet):
    field, value = EXAMPLE_BY_SHEET[sheet]
    return str(record.get(field, "")).strip() == value


def yes(value):
    return str(value).strip().lower() in {"yes", "y", "true", "x", "✓"}


def import_publications(ws, overwrite):
    n = 0
    for rec in rows(ws):
        if is_example(rec, "Publications"):
            continue
        title = str(rec.get("Title", "")).strip()
        if not title:
            continue
        year = str(rec.get("Year", "")).strip()
        year_num = re.sub(r"[^\d]", "", year) or "0000"
        fm = [
            "---",
            f"title: {q(title)}",
            f"authors: {yaml_list(rec.get('Authors'))}",
            f"venue: {q(rec.get('Venue'))}",
            f"year: {year_num}",
            f"type: {q(str(rec.get('Type','')).strip().lower() or 'journal')}",
            f"topic: {q(rec.get('Project'))}",
            f"selected: {'true' if yes(rec.get('Selected?')) else 'false'}",
        ]
        doi = str(rec.get("DOI", "")).strip()
        pdf = str(rec.get("PDF link", "")).strip()
        if doi:
            fm.append(f"doi: {q(doi)}")
        if pdf:
            fm.append(f"pdf: {q(pdf)}")
        fm.append("---\n")
        name = f"{year_num}-{slugify(title)[:50]}.md"
        n += write(CONTENT / "publications" / name, "\n".join(fm), overwrite)
    return n


def import_alumni(ws, overwrite):
    n = 0
    for rec in rows(ws):
        if is_example(rec, "Alumni"):
            continue
        name = str(rec.get("Name", "")).strip()
        if not name:
            continue
        body = "\n".join([
            "---",
            f"name: {q(name)}",
            f"title: {q(rec.get('Role / Degree'))}",
            f"years: {q(rec.get('Years in lab'))}",
            f"now: {q(rec.get('Current position'))}",
            "---\n",
        ])
        n += write(CONTENT / "alumni" / f"{slugify(name)}.md", body, overwrite)
    return n


def import_news(ws, overwrite):
    n = 0
    for rec in rows(ws):
        if is_example(rec, "News"):
            continue
        headline = str(rec.get("Headline", "")).strip()
        if not headline:
            continue
        date = str(rec.get("Date", "")).strip()[:10]
        body = "\n".join([
            "---",
            f"title: {q(headline)}",
            f"date: {date}",
            f"summary: {q(rec.get('Summary'))}",
            f"tags: {yaml_list(rec.get('Tags'))}",
            "---\n",
            str(rec.get("Summary", "")).strip(),
            "",
        ])
        n += write(CONTENT / "news" / f"{slugify(headline)[:50]}.md", body, overwrite)
    return n


def import_projects(ws, overwrite):
    n = 0
    order = 1
    for rec in rows(ws):
        title = str(rec.get("Title", "")).strip()
        if not title or is_example(rec, "Projects"):
            continue
        body = "\n".join([
            "---",
            f"title: {q(title)}",
            f"summary: {q(rec.get('Short summary'))}",
            f"order: {order}",
            "---\n",
            str(rec.get("Full description", "")).strip(),
            "",
        ])
        n += write(CONTENT / "projects" / f"{slugify(title)}.md", body, overwrite)
        order += 1
    return n


IMPORTERS = {
    "Publications": import_publications,
    "Alumni": import_alumni,
    "News": import_news,
    "Projects": import_projects,
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", help="Path to the filled-in .xlsx")
    ap.add_argument("--overwrite", action="store_true",
                    help="Replace files that already exist")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    total = 0
    for sheet, importer in IMPORTERS.items():
        if sheet not in wb.sheetnames:
            print(f"(no '{sheet}' sheet — skipped)")
            continue
        print(f"\n{sheet}:")
        total += importer(wb[sheet], args.overwrite)
    print(f"\nDone. {total} file(s) written.")
    if total:
        print("Next: run 'npm run build' to check everything validates, then commit.")


if __name__ == "__main__":
    main()
